"""Campaign calendar and instrument status, read live from the shared workbook
``SAPHFIRE 2026 - Experiments and Instrument Status.xlsx`` (repo root).

Row 1 holds the dates, row 2 the experiment titles (a title merged across
several date columns is one multi-day experiment), row 3 the logbook (UTC; the
roof times match the SAPHIR roof signal), rows 4+ one instrument each with a
status note per day. The cell colour carries the status:

    green  -> running        yellow -> issue / partial     orange -> not running
    grey   -> no experiment  no fill -> not recorded

    python campaign.py       # print the experiments and the instrument matrix
"""
from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

import openpyxl

WORKBOOK = Path(__file__).resolve().parents[1] / "SAPHFIRE 2026 - Experiments and Instrument Status.xlsx"

# Titles that mark a day without an experiment (matched case-insensitively).
NON_EXPERIMENT = ("integration", "downday", "no experiment")

# Short names for the archive folders, by start date (agreed 2026-10-03).
# Folder = <start date>_<E-number>_<short name>, e.g. 2026-06-30_E01_Smoldering-Fire_Night-to-Day.
SHORT_NAMES = {
    "2026-06-29": "Integration",
    "2026-06-30": "Smoldering-Fire_Night-to-Day",
    "2026-07-02": "Smoldering-Fire_Day",
    "2026-07-03": "Flaming-Fire_Day",
    "2026-07-04": "Background",
    "2026-07-05": "Wall-Losses",
    "2026-07-06": "Flaming-Fire_Day",
    "2026-07-07": "Biomass-Burning_Night",
    "2026-07-08": "Calibrations",
    "2026-07-09": "Biogenic-Smoldering-Flaming_Day-Night-Day",
    "2026-07-11": "Background",
    "2026-07-12": "No-Experiment",
    "2026-07-13": "Furanoid-Phenolic-HC-Mix_Day",
    "2026-07-14": "Calibrations",
    "2026-07-15": "Smoldering-Fire_Day",
    "2026-07-16": "Isoprene-Flaming-NOx-Limonene_Day-Night-Day",
}
NON_EXPERIMENT_FOLDER = "Non-experiment days"

# Instrument rows that are not expected to deliver data at all (GC never ran).
EXCLUDED_INSTRUMENTS = ("GC-MS",)

# Fill colour -> status. Theme colours are (theme index, tint) as openpyxl reports them.
STATUS_BY_FILL = {
    (9, 0.8): "running",          # accent6 green, light
    "FFFFEB9C": "issue",          # Excel 'Neutral' yellow
    (5, 0.8): "not running",      # accent2 orange, light
    (0, -0.05): "no experiment",  # background grey
}
STATUS_SYMBOL = {"running": "+", "issue": "~", "not running": "x", "no experiment": " ",
                 "unknown": "?"}


@dataclass
class Experiment:
    start: dt.date
    days: list[dt.date]
    title: str
    is_experiment: bool
    number: int | None = None          # E01... in date order, experiments only
    logbook: dict[dt.date, str] = field(default_factory=dict)

    @property
    def label(self) -> str:
        return f"E{self.number:02d}" if self.number else "--"

    @property
    def short_name(self) -> str:
        name = SHORT_NAMES.get(str(self.start))
        if name:
            return name
        return re.sub(r"[^A-Za-z0-9]+", "-", self.title).strip("-")[:40] or "Day"

    @property
    def folder(self) -> str:
        """Archive folder relative to Archive - Download."""
        if self.is_experiment:
            return f"{self.start}_{self.label}_{self.short_name}"
        return f"{NON_EXPERIMENT_FOLDER}/{self.start}_{self.short_name}"


def _fill_status(cell) -> str:
    if not cell.fill or cell.fill.fill_type is None:
        return "unknown"
    c = cell.fill.fgColor
    key = c.rgb if c.type == "rgb" else (c.theme, round(c.tint, 2))
    return STATUS_BY_FILL.get(key, "unknown")


@lru_cache(maxsize=1)
def _load():
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb.worksheets[0]
    cols = {}                                   # column index -> date
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        if isinstance(v, dt.datetime):
            cols[c] = v.date()
    spans = {}                                  # first column -> list of columns (merged titles)
    for rng in ws.merged_cells.ranges:
        if rng.min_row == 2 and rng.min_col in cols:
            spans[rng.min_col] = list(range(rng.min_col, rng.max_col + 1))
    covered = {c for cs in spans.values() for c in cs}

    exps, n = [], 0
    for c in sorted(cols):
        if c in covered and c not in spans:
            continue
        cs = spans.get(c, [c])
        title = " ".join(str(ws.cell(2, c).value or "").split())
        is_exp = bool(title) and not any(k in title.lower() for k in NON_EXPERIMENT)
        if is_exp:
            n += 1
        exps.append(Experiment(
            start=cols[cs[0]], days=[cols[x] for x in cs], title=title or "(no title)",
            is_experiment=is_exp, number=n if is_exp else None,
            logbook={cols[x]: str(ws.cell(3, x).value or "").strip() for x in cs}))

    status = {}                                 # instrument -> {date: (status, note)}
    for r in range(4, ws.max_row + 1):
        name = ws.cell(r, 2).value
        if not name:
            continue
        name = " ".join(str(name).split())
        if name in EXCLUDED_INSTRUMENTS:
            continue
        status[name] = {}
        for c, d in cols.items():
            note = str(ws.cell(r, c).value or "").strip()
            st = _fill_status(ws.cell(r, c))
            if note.lower().startswith("not running"):   # text wins over a yellow fill
                st = "not running"
            status[name][d] = (st, note)
    return exps, status


def experiments(only_experiments: bool = False) -> list[Experiment]:
    exps, _ = _load()
    return [e for e in exps if e.is_experiment or not only_experiments]


def campaign_days() -> list[dt.date]:
    return sorted(d for e in experiments() for d in e.days)


def experiment_for(day: dt.date) -> Experiment | None:
    return next((e for e in experiments() if day in e.days), None)


def instrument_status() -> dict[str, dict[dt.date, tuple[str, str]]]:
    return _load()[1]


def experiment_status(e: Experiment) -> dict[str, str]:
    """Worst status per instrument over the experiment's days."""
    order = ["not running", "issue", "unknown", "running", "no experiment"]
    out = {}
    for inst, days in instrument_status().items():
        sts = [days[d][0] for d in e.days if d in days]
        out[inst] = min(sts, key=order.index) if sts else "unknown"
    return out


if __name__ == "__main__":
    exps = experiments()
    insts = list(instrument_status())
    print(f"{'':4s} {'start':10s} {'days':4s} " + " ".join(f"{i[:4]:>4s}" for i in insts) + "  title")
    for e in exps:
        st = experiment_status(e)
        print(f"{e.label:4s} {e.start} {len(e.days):>4d} "
              + " ".join(f"{STATUS_SYMBOL[st[i]]:>4s}" for i in insts) + f"  {e.title[:60]}")
    print("\n+ running  ~ issue/partial  x not running  ? not recorded")
    print("Instruments:", "; ".join(f"{i[:4]}={i}" for i in insts))
